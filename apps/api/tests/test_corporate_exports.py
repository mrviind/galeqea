"""Corporate exports (Word/Excel) + page-aware chat (owner stabilization directive)."""

from __future__ import annotations

import io
import zipfile


def _run_with_results(db, project):
    from galeqea.models import Run, RunStatus, RunTest
    run = Run(project_id=project.id, number=7, title="Floor", status=RunStatus.FAILED,
              base_url="https://example.com")
    db.add(run)
    db.flush()
    db.add(RunTest(run_id=run.id, test_key="EX-GP-FUNC-01", title="Home loads",
                   status="passed", test_case_id="", duration_ms=800))
    db.add(RunTest(run_id=run.id, test_key="EX-GP-A11Y-01", title="a11y: /",
                   status="failed", test_case_id="", duration_ms=1200,
                   error_type="accessibility", error_message="1 serious (color-contrast)"))
    db.commit()
    return run


def _cases(db, project):
    from galeqea.models import TestCase, TestStep
    c = TestCase(project_id=project.id, key="EX-T-0001", title="verify login",
                 status="proposed", priority="high", risk="high",
                 requirement_refs=["REQ-001"], covers=["REQ-001-R1"],
                 provenance={"technique": "boundary_value"})
    db.add(c)
    db.flush()
    db.add(TestStep(test_case_id=c.id, index=0, action="fill", intent="enter password",
                    expected="accepted"))
    db.commit()
    return [c]


def test_test_cases_xlsx(db, project):
    from galeqea.reports.exports_xlsx import test_cases_workbook
    data = test_cases_workbook(_cases(db, project), project_name=project.name)
    assert data[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Test Cases"]
    ws = wb["Test Cases"]
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "EX-T-0001" in text and "verify login" in text and "boundary_value" in text


def test_tcr_xlsx_has_three_sheets(db, project):
    from galeqea.reports.exports_xlsx import tcr_workbook
    run = _run_with_results(db, project)
    data = tcr_workbook(db, project, run)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Results", "By Type"]
    results = "\n".join(str(c.value) for row in wb["Results"].iter_rows() for c in row if c.value)
    assert "color-contrast" in results


def test_tcr_xlsx_summary_has_a_doughnut_and_a_bar_chart(db, project):
    """The dashboard sheet: a native, editable doughnut (pass/fail/skip) and a
    native stacked bar (by test type), not a static image, so a stakeholder can
    click into the chart's own data in Excel."""
    from openpyxl.chart import BarChart, DoughnutChart

    from galeqea.reports.exports_xlsx import tcr_workbook
    run = _run_with_results(db, project)
    data = tcr_workbook(db, project, run)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    charts = wb["Summary"]._charts
    assert any(isinstance(c, DoughnutChart) for c in charts)
    assert any(isinstance(c, BarChart) for c in charts)
    # KPI cards: the big numbers are real cell values, not baked into an image,
    # so a reader (or a formula) can use them directly. _run_with_results makes
    # exactly 1 passed + 1 failed = 2 total.
    values = {c.value for row in wb["Summary"].iter_rows() for c in row}
    assert {"2", "1", "50%"} <= values


def test_tcr_xlsx_survives_control_characters_in_error_text(db, project):
    """Regression: a raw control character in an error message (a Playwright
    timeout message, terminal color codes, a stray NUL) used to make openpyxl
    raise IllegalCharacterError and 500 the whole export. common.strip_control_chars
    is applied in build_run_report; this proves the xlsx path survives it end to end."""
    from galeqea.models import Run, RunStatus, RunTest
    from galeqea.reports.exports_xlsx import tcr_workbook
    run = Run(project_id=project.id, number=8, title="Floor", status=RunStatus.FAILED,
              base_url="https://example.com")
    db.add(run)
    db.flush()
    db.add(RunTest(run_id=run.id, test_key="EX-CTRL-01", title="control char test",
                   status="failed", test_case_id="", duration_ms=100,
                   error_type="timeout", error_message="Timeout\x1b[31m exceeded\x00"))
    db.commit()
    data = tcr_workbook(db, project, run)  # must not raise
    assert data[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    text = "\n".join(str(c.value) for row in wb["Results"].iter_rows() for c in row if c.value)
    assert "Timeout" in text and "exceeded" in text


def test_rtm_xlsx(db, project):
    from galeqea.reports.exports_xlsx import rtm_workbook
    matrix = [{"ref": "REQ-001", "title": "Login", "risk": "high", "priority": "P1",
               "covered": False, "rules": [{"rule_id": "REQ-001-R1", "rule_type": "validation",
                                            "technique": "boundary_value", "tests": ["EX-T-0001"],
                                            "covered": False}]}]
    data = rtm_workbook(matrix, project_name=project.name)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    txt = "\n".join(str(c.value) for row in wb["Traceability"].iter_rows() for c in row if c.value)
    assert "REQ-001" in txt and "REQ-001-R1" in txt


def test_rtm_xlsx_has_a_dashboard_sheet_with_a_doughnut(db, project):
    from openpyxl.chart import DoughnutChart

    from galeqea.reports.exports_xlsx import rtm_workbook
    matrix = [{"ref": "REQ-001", "title": "Login", "risk": "high", "priority": "P1",
               "covered": True, "rules": []},
              {"ref": "REQ-002", "title": "Logout", "risk": "low", "priority": "P3",
               "covered": False, "rules": []}]
    data = rtm_workbook(matrix, project_name=project.name)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert "Dashboard" in wb.sheetnames
    assert wb.sheetnames[0] == "Dashboard"  # the sheet that opens first
    assert any(isinstance(c, DoughnutChart) for c in wb["Dashboard"]._charts)


def _docx_text(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def test_test_plan_docx(db, project):
    from galeqea.models import Journey
    from galeqea.reports.test_plan_docx import build_test_plan_docx
    jrn = Journey(project_id=project.id, target="https://example.com",
                  plan={"pages": ["https://example.com/", "https://example.com/about"],
                        "functional": ["Each page loads"], "non_functional": ["No 5xx"],
                        "typed": {"types": [{"key": "a11y", "label": "Accessibility",
                                             "enabled": True, "count": 2, "risk": "high"}],
                                  "totals": {"test_count": 2}}})
    db.add(jrn)
    db.commit()
    data = build_test_plan_docx(db, project, jrn)
    assert data[:2] == b"PK"
    text = _docx_text(data)
    for section in ["Test Plan", "1. Introduction", "2. Scope", "3. Features to be tested",
                    "5. Entry & exit criteria", "9. Approvals"]:
        assert section in text
    assert "HIGH" in text                          # the risk pill, upper-cased
    assert "Planned cases" in text or "2" in text   # the cover's stat card
    assert zipfile.is_zipfile(io.BytesIO(data))
    z = zipfile.ZipFile(io.BytesIO(data))
    assert 'TOC \\o "1-2" \\h \\z \\u' in z.read("word/document.xml").decode("utf-8")
    # Calibri, not Aptos - see docx_style.py's module docstring (a verified
    # cross-renderer consistency finding, not a stale default).
    assert "Calibri" in z.read("word/styles.xml").decode("utf-8")


# --- page-aware chat ------------------------------------------------------- #

def test_page_chat_answers_tests_page(db, project, humans):
    from galeqea.services import page_chat
    _cases(db, project)
    out = page_chat.try_handle(db, project, "how many test cases are there?",
                               humans["author"], "/tests")
    assert out and "test case" in out[0].lower()
    assert out[1][0]["type"] == "page_answer" and out[1][0]["page"] == "tests"


def test_page_chat_export_is_page_aware(db, project, humans):
    from galeqea.services import page_chat
    out = page_chat.try_handle(db, project, "export test cases as excel",
                               humans["author"], "/tests")
    assert out and out[1][0]["artifact"] == "test_cases" and out[1][0]["format"] == "xlsx"
    out = page_chat.try_handle(db, project, "export the test plan", humans["author"], "/tests")
    assert out and out[1][0]["artifact"] == "test_plan"


def test_page_chat_ignores_non_questions(db, project, humans):
    from galeqea.services import page_chat
    assert page_chat.try_handle(db, project, "run the smoke tests",
                                humans["author"], "/tests") is None
    assert page_chat.try_handle(db, project, "hello", humans["author"], "/tests") is None
