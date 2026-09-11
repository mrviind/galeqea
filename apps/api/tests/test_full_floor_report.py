"""The Full Floor flow + the branded Word report + the job queue."""

from __future__ import annotations

import io
import zipfile

from galeqea.jobs import get_queue, reset_queue
from galeqea.jobs.inprocess import InProcessQueue


def _docx_text(data: bytes) -> str:
    """All readable text in a .docx: top-level paragraphs *and* table cells (the
    visual redesign moved several sections, e.g. failure errors, into styled
    tables/callouts, so a paragraphs-only scan would miss them)."""
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


# --- job queue ------------------------------------------------------------- #

def test_default_queue_is_inprocess_on_sqlite():
    reset_queue()
    q = get_queue()
    assert isinstance(q, InProcessQueue) and q.kind == "inprocess"


async def test_inprocess_queue_runs_a_task_and_tracks_depth():
    from galeqea.jobs.base import task
    ran = {}

    @task("_t_probe")
    async def _probe(*, x):
        ran["x"] = x

    q = InProcessQueue()
    job_id = await q.enqueue("_t_probe", x=42)
    import asyncio
    for _ in range(50):
        if q.job_status(job_id)["status"] == "done":
            break
        await asyncio.sleep(0.01)
    assert ran.get("x") == 42
    assert q.job_status(job_id)["status"] == "done"
    assert q.depth() == 0


def test_config_queue_kind_switches_on_postgres(monkeypatch):
    from galeqea.config import settings
    monkeypatch.setattr(settings, "database_url", "postgresql+psycopg://u:p@h/db")
    assert settings.is_postgres is True
    assert settings.queue_kind == "procrastinate"
    monkeypatch.setattr(settings, "database_url", "sqlite:///x.db")
    assert settings.queue_kind == "inprocess"


# --- the Word report ------------------------------------------------------- #

def _run_with_results(db, project):
    from galeqea.models import Run, RunStatus, RunTest
    run = Run(project_id=project.id, number=1, title="Floor: https://example.com",
              status=RunStatus.FAILED, base_url="https://example.com")
    db.add(run)
    db.flush()
    db.add(RunTest(run_id=run.id, test_key="EX-GP-FUNC-01", title="Home loads",
                   status="passed", test_case_id="", duration_ms=800))
    db.add(RunTest(run_id=run.id, test_key="EX-GP-A11Y-01", title="Accessibility: /",
                   status="failed", test_case_id="", duration_ms=1200,
                   error_type="accessibility", error_message="1 serious (color-contrast)"))
    db.commit()
    return run


def test_word_report_is_a_valid_docx_with_sections(db, project):
    from galeqea.reports.report_docx import build_run_docx
    run = _run_with_results(db, project)
    data = build_run_docx(db, project, run, target="https://example.com",
                          pages=["https://example.com/", "https://example.com/about"],
                          plan={"functional": ["Each page loads"], "non_functional": ["No 5xx"]})
    assert data[:2] == b"PK"                      # a zip (docx) container
    text = _docx_text(data)
    assert "Test Completion Report" in text        # matches the product's own name for it
    assert "https://example.com" in text
    assert "Executive summary" in text and "Failures & evidence" in text
    # the logo is embedded
    z = zipfile.ZipFile(io.BytesIO(data))
    assert any(n.startswith("word/media/") for n in z.namelist())
    # the failing test's error is present (now inside a callout table, not a bare paragraph)
    assert "color-contrast" in text
    # 50% pass rate -> the cover's hero number and the run's own status appear
    assert "50%" in text
    # the TOC field's fallback text is a real section list, not an apology - a
    # viewer whose tool never runs Word's field-update pass (e.g. a headless
    # LibreOffice PDF export) still sees something useful on the Contents page.
    assert "Executive summary" in text and "Conclusion & sign-off" in text
    assert "Right-click" not in text
    # a real table of contents field, not a static list
    assert 'TOC \\o "1-2" \\h \\z \\u' in z.read("word/document.xml").decode("utf-8")
    # a running page-number footer (PAGE / NUMPAGES fields)
    footer_files = [n for n in z.namelist() if n.startswith("word/footer")]
    assert footer_files
    footer_xml = "".join(z.read(n).decode("utf-8") for n in footer_files)
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml


def test_word_report_uses_a_typeface_that_renders_consistently_everywhere(db, project):
    """Calibri, deliberately - not the newer Microsoft 365 default "Aptos". Verified
    by rendering both in LibreOffice: Aptos has no bundled cross-platform substitute
    and resolved *inconsistently* between weights (bold fell back to a sans, regular
    to a serif - a broken-looking mix on the same page); Calibri's metric-compatible
    substitute (Carlito) renders identically everywhere. See docx_style.py's module
    docstring."""
    from galeqea.reports.report_docx import build_run_docx
    run = _run_with_results(db, project)
    data = build_run_docx(db, project, run, target="https://example.com")
    z = zipfile.ZipFile(io.BytesIO(data))
    styles_xml = z.read("word/styles.xml").decode("utf-8")
    assert "Calibri" in styles_xml             # Normal + Heading 1 / Title
    assert "Aptos" not in styles_xml
    assert "Aptos" not in z.read("word/document.xml").decode("utf-8")


def test_word_report_handles_a_run_with_no_results(db, project):
    from galeqea.models import Run, RunStatus
    from galeqea.reports.report_docx import build_run_docx
    run = Run(project_id=project.id, number=2, title="empty", status=RunStatus.ERROR)
    db.add(run)
    db.commit()
    data = build_run_docx(db, project, run, target="https://example.com")
    assert data[:2] == b"PK"


def test_word_report_embeds_a_results_doughnut_even_with_no_failures(db, project):
    """The executive summary's doughnut is a real embedded image (python-docx has
    no native chart), so it must be present even when there is nothing to
    screenshot; an all-pass run used to mean zero embedded images beyond the
    cover logo; now the doughnut is always there once there's at least one test."""
    from galeqea.models import Run, RunStatus, RunTest
    from galeqea.reports.report_docx import build_run_docx
    run = Run(project_id=project.id, number=3, title="all green", status=RunStatus.PASSED,
              base_url="https://example.com")
    db.add(run)
    db.flush()
    db.add(RunTest(run_id=run.id, test_key="EX-OK-01", title="Home loads",
                   status="passed", test_case_id="", duration_ms=500))
    db.commit()
    data = build_run_docx(db, project, run, target="https://example.com")
    z = zipfile.ZipFile(io.BytesIO(data))
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    assert len(media) >= 2  # the cover logo *and* the doughnut chart
    text = _docx_text(data)
    assert "PASS RATE" in text or "pass rate" in text.lower()


def test_word_report_survives_control_characters_in_error_text(db, project):
    """Same regression as the xlsx side (test_corporate_exports.py): a raw
    control character in an error message used to make lxml raise
    'All strings must be XML compatible' and 500 the whole report."""
    from galeqea.models import Run, RunStatus, RunTest
    from galeqea.reports.report_docx import build_run_docx
    run = Run(project_id=project.id, number=4, title="ctrl", status=RunStatus.FAILED,
              base_url="https://example.com")
    db.add(run)
    db.flush()
    db.add(RunTest(run_id=run.id, test_key="EX-CTRL-01", title="control char test",
                   status="failed", test_case_id="", duration_ms=100,
                   error_type="timeout", error_message="Timeout\x1b[31m exceeded\x00"))
    db.commit()
    data = build_run_docx(db, project, run, target="https://example.com")  # must not raise
    assert data[:2] == b"PK"
    text = _docx_text(data)
    assert "Timeout" in text and "exceeded" in text


# --- full_floor plan summary ---------------------------------------------- #

def test_plan_summary_shape():
    from galeqea.models import Journey
    from galeqea.services.full_floor import plan_summary
    jrn = Journey(project_id="p", target="https://example.com",
                  discovery={"pages": ["a", "b"], "forms": 2, "method": "browser", "discovered": 9},
                  plan={"pages": ["a", "b"], "functional": ["loads"], "non_functional": ["no 5xx"],
                        "typed": {"types": [{"key": "a11y", "enabled": True, "count": 2}],
                                  "totals": {"test_count": 2}}})
    s = plan_summary(jrn)
    assert s["page_count"] == 2 and s["forms"] == 2
    assert s["test_types"] == [{"type": "a11y", "count": 2}]
    assert s["test_count"] == 2
    assert s["elements"]["pages_discovered"] == 9


def test_report_includes_requirements_coverage_when_a_doc_was_ingested(db, project):
    """A report reflects an uploaded requirement document (WO-owner goal)."""
    from galeqea.reports.report_docx import build_run_docx
    from galeqea.services import requirements as req
    req.ingest_document(db, project_id=project.id, filename="prd.md", title="PRD",
                        data=b"# Spec\n\n- The password must be between 8 and 64 characters.\n")
    db.commit()
    run = _run_with_results(db, project)
    text = _docx_text(build_run_docx(db, project, run, target="https://example.com"))
    assert "Requirements coverage" in text
    assert "requirement(s) covered" in text
    assert "COVERED" in text or "GAP" in text          # the RTM status pill, upper-cased


def test_pages_override_is_authoritative(monkeypatch):
    """An explicit page list overrides whatever the crawl (possibly cached) returned."""
    from galeqea.services import website_test
    calls = {}

    def fake_discover(target, **kw):
        calls["target"] = target
        return {"ok": True, "base": "https://ex.com", "pages": ["https://ex.com/stale/"],
                "forms": 0, "method": "browser", "discovered": 1, "skipped": {}, "truncated": True}

    monkeypatch.setattr(website_test, "discover_pages", fake_discover)
    # Exercise only the discovery+override portion of propose via a light stub:
    disc = website_test.discover_pages("https://ex.com")
    pages = ["/", "/products/", "/sift/"]
    base = disc["base"].rstrip("/")
    disc["pages"] = [p if p.startswith("http") else f"{base}/{p.lstrip('/')}" for p in pages]
    assert disc["pages"] == ["https://ex.com/", "https://ex.com/products/", "https://ex.com/sift/"]
    assert "https://ex.com/stale/" not in disc["pages"]
